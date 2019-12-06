from docx import Document
import xlsxwriter

workbook = xlsxwriter.Workbook("TablaDeCuestionariosCentros_de_internamiento para Adolescentes.xlsx")
worksheet = workbook.add_worksheet()

def EscribirTitulos():
    
    worksheet.write(0,0,"GEN")
    worksheet.write(1,0,"M")
    worksheet.write(2,0,"H")
    worksheet.write(3,0,"M")
    worksheet.write(4,0,"H")

    worksheet.write(0,1,"FECHA")
    worksheet.write(1,1,"2018")
    worksheet.write(2,1,"2018")
    worksheet.write(3,1,"2019")
    worksheet.write(4,1,"2019")
    



#SE INICIALIZA LA VARIABLE CABECERA. SE USARÁ PARA SEPARAR LAS PREGUNTAS DE LOS TITULOS
cabeceras = []

#SE CREA UNA LISTA FINAL DONDE SE GUARDARÁ CADA PREGUNTA CON SUS RESPECTIVOS RESULTADOS
listaF = []


#============================SE ABRE EL DOCUMENTO INDICADO============================
wordDoc = Document('Centros_de_internamiento para Adolescentes.docx')

#====================SE ITERA SOBRE CADA TABLA, CADA RENGLÓN Y CADA CELDA DEL RENGLÓN====================
for table in wordDoc.tables:
    for row in table.rows:
        #===SE CREA UN INDICE PARA SABER EN QUE CELDA ESTAMOS (LAS PREGUNTAS SIEMPRE OCUPAN DOS CELDAS)===
        indice = 0
        respuestaTemp = []
        rPreg = ""
        rRes = []        
        for cell in row.cells:             
            #SI LA CELDA UNO DEL RENGLÓN ESTA VACÍA O SON SOLO MAYÚSCULAS SE TOMARÁ COMO CABECERA
            if row.cells[0].text.isupper() or row.cells[0].text == "":                
                cabeceras.append(cell.text)

            else:
                if indice < 2 :
                    #SI EL TEXTO SE ENCUENTRA DENTRO DE LAS PRIMERAS DOS COLUMNAS SE TOMARÁ COMO PREGUNTA
                    rPreg = row.cells[0].text
                    indice = indice + 1

                else:
                    #EN CASO CONTRARIO SE AGREGA EL TEXTO A LA LISTA DE RESPUESTAS
                    rRes.append(cell.text)                    

        #AL FINAL DE CADA RENGLÓN SE AGREGARÁN A LA LISTA FINAL TANTO LA PREGUNTA COMO SUS RESPECTIVAS RESPUESTAS
        if rPreg != "":
            listaF.append({"Pregunta":rPreg,"Respuestas":rRes})

#=================SE IMPRIME SOLO PARA COMPROBACIÓN Y PARA DEBUGGEO=================
print(listaF)

#INDICE
i=0

for l in listaF:
    #POR CADA ELEMENTO EN LA LISTA FINAL, SE ESCRIBE LA PREGUNTA EN LA CELDA ADECUADA
    worksheet.write(0, i+2, l["Pregunta"])
    #========== SE ITERA SOBRE CADA RESPUESTA QUE CONTENGA LA LISTA ==========
    for j in range(len(l["Respuestas"])):
        worksheet.write(j+1, i+2, l ["Respuestas"][j] )  

    i = i+1
     

EscribirTitulos()


print("Excel file ready") 

workbook.close()